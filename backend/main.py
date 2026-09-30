import asyncio
import base64
import binascii
import importlib.util
import json
import logging
import re
import time
import uuid
import zipfile
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path, PurePath
from typing import Annotated

from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer

from backend.auth import AuthService, User
from backend.chats import ChatStore
from backend.config import Settings, get_settings
from backend.email_service import EmailConfigurationError, send_password_reset_email
from backend.rag import (
    MediaGenerationError,
    RAGConfigurationError,
    RAGService,
    RAGServiceError,
)
from backend.rate_limit import InMemoryRateLimiter
from backend.schemas import (
    AttachmentPreviewResponse,
    AuthResponse,
    BoardCreateRequest,
    BoardDetailResponse,
    BoardRenameRequest,
    BoardSummaryResponse,
    ChatAttachment,
    ChatRequest,
    ChatResponse,
    ForgotPasswordRequest,
    HealthResponse,
    LoginRequest,
    MessageResponseBody,
    RegisterRequest,
    ResetPasswordRequest,
    SampleQuestionsResponse,
    SearchResultResponse,
    TemporaryUploadResponse,
    UserResponse,
)


@asynccontextmanager
async def lifespan(_: FastAPI) -> AsyncIterator[None]:
    settings.validate_security()
    if MISSING_ATTACHMENT_DEPENDENCIES:
        raise RuntimeError(
            "Missing attachment dependencies: "
            + ", ".join(MISSING_ATTACHMENT_DEPENDENCIES)
            + '. Run: python -m pip install -e ".[dev]"'
        )
    get_chat_store(settings).cleanup_temporary_uploads()
    cleanup_task = asyncio.create_task(_cleanup_expired_uploads())
    try:
        yield
    finally:
        cleanup_task.cancel()
        try:
            await cleanup_task
        except asyncio.CancelledError:
            pass


app = FastAPI(
    title="FlowDesk Support API",
    description="Grounded customer support using OpenAI File Search.",
    version="1.0.0",
    lifespan=lifespan,
)
logger = logging.getLogger(__name__)
rate_limiter = InMemoryRateLimiter()

settings = get_settings()
app.add_middleware(
    CORSMiddleware,
    allow_origins=[settings.frontend_origin],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Content-Type", "Authorization"],
)

SAMPLE_QUESTIONS = [
    "How do I invite a teammate?",
    "What happens when I cancel my subscription?",
    "How can I fix error FD-403?",
]

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_TOTAL_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_BINARY_UPLOAD_BYTES = 100 * 1024 * 1024
MAX_EXTRACTED_CHARS = 50_000
SUPPORTED_ATTACHMENT_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
TEXT_ATTACHMENT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}
IMAGE_ATTACHMENT_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
ATTACHMENT_DEPENDENCIES = {
    "pypdf": "pypdf",
    "python-docx": "docx",
    "openpyxl": "openpyxl",
    "python-pptx": "pptx",
    "Pillow": "PIL",
}
MISSING_ATTACHMENT_DEPENDENCIES = [
    package
    for package, module in ATTACHMENT_DEPENDENCIES.items()
    if importlib.util.find_spec(module) is None
]
MAX_OFFICE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_OFFICE_ARCHIVE_ENTRIES = 10_000
if MISSING_ATTACHMENT_DEPENDENCIES:
    logger.warning(
        "Attachment analysis dependencies are missing: %s. Run: python -m pip install -e .",
        ", ".join(MISSING_ATTACHMENT_DEPENDENCIES),
    )


class AttachmentValidationError(ValueError):
    pass


def _validate_office_archive(path: Path, filename: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_OFFICE_ARCHIVE_ENTRIES:
                raise AttachmentValidationError(f"{filename} contains too many archive entries.")
            expanded = sum(member.file_size for member in members)
            if expanded > MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise AttachmentValidationError(
                    f"{filename} expands beyond the safe 200 MB analysis limit."
                )
            if any(member.file_size > MAX_OFFICE_UNCOMPRESSED_BYTES for member in members):
                raise AttachmentValidationError(f"{filename} contains an unsafe archive entry.")
    except zipfile.BadZipFile as exc:
        raise AttachmentValidationError(f"{filename} is corrupted or mislabeled.") from exc


def _attachment_text(text: str, filename: str) -> str:
    cleaned = text.strip()
    if not cleaned:
        raise AttachmentValidationError(f"{filename} does not contain readable text.")
    if len(cleaned) > MAX_EXTRACTED_CHARS:
        return cleaned[:MAX_EXTRACTED_CHARS] + "\n[Content truncated at 50,000 characters]"
    return cleaned


def _extract_office_text(extension: str, data: bytes, filename: str) -> str:
    if extension == ".docx":
        from docx import Document

        document = Document(BytesIO(data))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"Worksheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                if any(values):
                    lines.append(" | ".join(values))
                if sum(map(len, lines)) >= MAX_EXTRACTED_CHARS:
                    break
        workbook.close()
    else:
        from pptx import Presentation

        lines = []
        for index, slide in enumerate(Presentation(BytesIO(data)).slides, start=1):
            lines.append(f"Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text)
                if getattr(shape, "has_table", False):
                    lines.extend(
                        " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                    )
    return _attachment_text("\n".join(lines), filename)


def prepare_attachments(attachments: list[ChatAttachment]) -> list[dict]:
    parts: list[dict] = []
    total = 0
    for attachment in attachments:
        if (
            PurePath(attachment.filename).name != attachment.filename
            or "\x00" in attachment.filename
        ):
            raise AttachmentValidationError("Attachment filenames cannot contain a path.")
        extension = PurePath(attachment.filename).suffix.lower()
        expected_mime = SUPPORTED_ATTACHMENT_TYPES.get(extension)
        if not expected_mime:
            raise AttachmentValidationError(
                f"{attachment.filename} has an unsupported file format."
            )
        expected_encoding = "utf8" if extension in TEXT_ATTACHMENT_EXTENSIONS else "base64"
        if attachment.mime_type != expected_mime or attachment.encoding != expected_encoding:
            raise AttachmentValidationError(f"{attachment.filename} has an invalid file type.")
        try:
            data = (
                attachment.content.encode("utf-8")
                if attachment.encoding == "utf8"
                else base64.b64decode(attachment.content, validate=True)
            )
        except (UnicodeError, binascii.Error, ValueError) as exc:
            raise AttachmentValidationError(
                f"{attachment.filename} contains invalid file data."
            ) from exc
        if not data:
            raise AttachmentValidationError(f"{attachment.filename} is empty.")
        if len(data) != attachment.size:
            raise AttachmentValidationError(f"{attachment.filename} has an invalid reported size.")
        total += len(data)
        if len(data) > MAX_ATTACHMENT_BYTES or total > MAX_TOTAL_ATTACHMENT_BYTES:
            raise AttachmentValidationError("Attachments exceed the allowed size limit.")
        signatures = {
            ".pdf": data.startswith(b"%PDF-"),
            ".jpg": data.startswith(b"\xff\xd8\xff"),
            ".jpeg": data.startswith(b"\xff\xd8\xff"),
            ".png": data.startswith(b"\x89PNG\r\n\x1a\n"),
            ".webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP",
            ".docx": data.startswith(b"PK"),
            ".xlsx": data.startswith(b"PK"),
            ".pptx": data.startswith(b"PK"),
        }
        if extension in signatures and not signatures[extension]:
            raise AttachmentValidationError(f"{attachment.filename} is corrupted or mislabeled.")
        if extension in TEXT_ATTACHMENT_EXTENSIONS:
            text = _attachment_text(data.decode("utf-8-sig"), attachment.filename)
            parts.append(
                {"type": "input_text", "text": f"Attached file: {attachment.filename}\n{text}"}
            )
        elif extension == ".pdf":
            try:
                from pypdf import PdfReader

                if PdfReader(BytesIO(data)).is_encrypted:
                    raise AttachmentValidationError(f"{attachment.filename} is password-protected.")
            except AttachmentValidationError:
                raise
            except ModuleNotFoundError as exc:
                raise AttachmentValidationError(
                    "PDF analysis is not installed. Run 'python -m pip install -e .' "
                    "and restart the backend."
                ) from exc
            except Exception as exc:
                raise AttachmentValidationError(
                    f"{attachment.filename} is not a readable PDF."
                ) from exc
            parts.append(
                {
                    "type": "input_file",
                    "filename": attachment.filename,
                    "file_data": f"data:application/pdf;base64,{attachment.content}",
                }
            )
        elif extension in IMAGE_ATTACHMENT_EXTENSIONS:
            parts.append(
                {
                    "type": "input_image",
                    "image_url": f"data:{attachment.mime_type};base64,{attachment.content}",
                    "detail": "auto",
                }
            )
        else:
            try:
                text = _extract_office_text(extension, data, attachment.filename)
            except AttachmentValidationError:
                raise
            except ModuleNotFoundError as exc:
                raise AttachmentValidationError(
                    "Office-file analysis is not installed. Run 'python -m pip install -e .' "
                    "and restart the backend."
                ) from exc
            except Exception as exc:
                raise AttachmentValidationError(
                    f"{attachment.filename} is corrupted or unreadable."
                ) from exc
            parts.append(
                {"type": "input_text", "text": f"Extracted from {attachment.filename}:\n{text}"}
            )
    return parts


def prepare_attachment_storage(attachments: list[ChatAttachment]) -> list[dict]:
    stored: list[dict] = []
    for attachment in attachments:
        extension = PurePath(attachment.filename).suffix.lower()
        data = (
            attachment.content.encode("utf-8")
            if attachment.encoding == "utf8"
            else base64.b64decode(attachment.content, validate=True)
        )
        preview_text: str | None = None
        if extension in TEXT_ATTACHMENT_EXTENSIONS:
            preview_type = "text"
            preview_text = _attachment_text(data.decode("utf-8-sig"), attachment.filename)
        elif extension == ".pdf":
            preview_type = "pdf"
        elif extension in IMAGE_ATTACHMENT_EXTENSIONS:
            preview_type = "image"
        else:
            preview_type = "document"
            preview_text = _extract_office_text(extension, data, attachment.filename)
        stored.append(
            {
                "filename": attachment.filename,
                "mime_type": attachment.mime_type,
                "data": data,
                "preview_type": preview_type,
                "preview_text": preview_text,
            }
        )
    return stored


def prepare_temporary_uploads(uploads: list[dict]) -> tuple[list[dict], list[dict]]:
    parts: list[dict] = []
    prepared: list[dict] = []
    for upload in uploads:
        path: Path = upload["path"]
        filename = upload["filename"]
        extension = path.suffix.lower()
        preview_text: str | None = None
        if extension in TEXT_ATTACHMENT_EXTENSIONS:
            text = _attachment_text(
                path.read_text(encoding="utf-8-sig")[: MAX_EXTRACTED_CHARS + 1], filename
            )
            preview_text = text
            parts.append({"type": "input_text", "text": f"Attached file: {filename}\n{text}"})
        elif extension == ".pdf":
            from pypdf import PdfReader

            if PdfReader(path).is_encrypted:
                raise AttachmentValidationError(f"{filename} is password-protected.")
            parts.append({"type": "local_input_file", "filename": filename, "path": str(path)})
        elif extension in IMAGE_ATTACHMENT_EXTENSIONS:
            from PIL import Image

            with Image.open(path) as image:
                image.thumbnail((2048, 2048))
                output = BytesIO()
                image_format = "PNG" if extension == ".png" else "JPEG"
                if image_format == "JPEG" and image.mode not in {"RGB", "L"}:
                    image = image.convert("RGB")
                image.save(output, format=image_format, optimize=True, quality=88)
            mime_type = "image/png" if image_format == "PNG" else "image/jpeg"
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
            parts.append(
                {
                    "type": "input_image",
                    "image_url": f"data:{mime_type};base64,{encoded}",
                    "detail": "auto",
                }
            )
        else:
            preview_text = _extract_office_text(extension, path.read_bytes(), filename)
            parts.append(
                {"type": "input_text", "text": f"Extracted from {filename}:\n{preview_text}"}
            )
        prepared.append({**upload, "preview_text": preview_text})
    return parts, prepared


SettingsDependency = Annotated[Settings, Depends(get_settings)]
bearer = HTTPBearer(auto_error=False)


def get_rag_service(config: SettingsDependency) -> RAGService:
    return RAGService(
        api_key=config.openai_api_key,
        vector_store_id=config.vector_store_id,
        model=config.openai_model,
    )


def get_auth_service(config: SettingsDependency) -> AuthService:
    return AuthService(
        database_path=config.auth_database_path,
        secret=config.auth_secret,
        token_minutes=config.auth_token_minutes,
    )


def get_chat_store(config: SettingsDependency) -> ChatStore:
    return ChatStore(
        config.auth_database_path,
        config.generated_media_path,
        config.uploaded_attachments_path,
        config.temporary_uploads_path,
    )


AuthDependency = Annotated[AuthService, Depends(get_auth_service)]
ChatStoreDependency = Annotated[ChatStore, Depends(get_chat_store)]


def get_current_user(
    credentials: Annotated[HTTPAuthorizationCredentials | None, Depends(bearer)],
    auth: AuthDependency,
) -> User:
    user = auth.user_from_token(credentials.credentials) if credentials else None
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    return user


CurrentUser = Annotated[User, Depends(get_current_user)]


def resolve_request_attachments(
    request: ChatRequest, chats: ChatStore, user_id: int
) -> tuple[list[dict], list[dict], list[dict]]:
    if request.attachments and request.attachment_ids:
        raise AttachmentValidationError(
            "Use uploaded attachment IDs or legacy attachments, not both."
        )
    legacy_parts = prepare_attachments(request.attachments)
    legacy_storage = prepare_attachment_storage(request.attachments)
    pending = (
        chats.get_temporary_uploads(user_id, request.attachment_ids)
        if request.attachment_ids
        else []
    )
    if sum(item["size"] for item in pending) > MAX_BINARY_UPLOAD_BYTES:
        raise AttachmentValidationError("Attachments cannot exceed 100 MB combined.")
    pending_parts, prepared_pending = prepare_temporary_uploads(pending)
    return [*legacy_parts, *pending_parts], legacy_storage, prepared_pending


IMAGE_TERM_PATTERN = re.compile(
    r"\b(?:images?|pictures?|photos?|pics?|illustrations?|logos?)\b", re.IGNORECASE
)
VIDEO_TERM_PATTERN = re.compile(r"\b(?:videos?|movies?|clips?|animations?)\b", re.IGNORECASE)
MEDIA_REQUEST_CUE_PATTERN = re.compile(
    r"\b(?:create|generate|make|draw|design|show|provide|give|send|find|get|want|need)\b"
    r"|\b(?:would like|looking for)\b",
    re.IGNORECASE,
)
INFORMATIONAL_MEDIA_PATTERN = re.compile(
    r"\b(?:explain|what|why|how|difference|compare|describe|analyze)\b.{0,60}"
    r"\b(?:images?|pictures?|photos?|pics?|illustrations?|logos?|videos?|movies?|clips?|animations?)\b",
    re.IGNORECASE,
)


def is_image_request(question: str) -> bool:
    return _is_media_request(question, IMAGE_TERM_PATTERN)


def is_video_request(question: str) -> bool:
    return _is_media_request(question, VIDEO_TERM_PATTERN)


def _is_media_request(question: str, term_pattern: re.Pattern[str]) -> bool:
    if not term_pattern.search(question):
        return False
    if INFORMATIONAL_MEDIA_PATTERN.search(question):
        return False
    return bool(MEDIA_REQUEST_CUE_PATTERN.search(question)) or len(question.split()) <= 12


async def _cleanup_expired_uploads() -> None:
    while True:
        await asyncio.sleep(max(60, settings.temporary_upload_cleanup_minutes * 60))
        try:
            get_chat_store(settings).cleanup_temporary_uploads()
        except Exception:
            logger.exception("Scheduled temporary-upload cleanup failed")


@app.get("/api/health", response_model=HealthResponse)
async def health(config: SettingsDependency) -> HealthResponse:
    return HealthResponse(
        status="ok" if config.configured else "setup_required",
        configured=config.configured,
        model=config.openai_model,
        attachment_support=not MISSING_ATTACHMENT_DEPENDENCIES,
        missing_attachment_dependencies=MISSING_ATTACHMENT_DEPENDENCIES,
    )


@app.get("/api/sample-questions", response_model=SampleQuestionsResponse)
async def sample_questions() -> SampleQuestionsResponse:
    return SampleQuestionsResponse(questions=SAMPLE_QUESTIONS)


@app.post("/api/auth/register", response_model=AuthResponse, status_code=201)
async def register(request: RegisterRequest, auth: AuthDependency) -> AuthResponse:
    rate_limiter.check(
        f"auth:register:{request.email.lower()}", settings.auth_rate_limit_per_minute
    )
    user = auth.register(request.email, request.name.strip(), request.password)
    return AuthResponse(access_token=auth.create_token(user), user=UserResponse(**user.__dict__))


@app.post("/api/auth/login", response_model=AuthResponse)
async def login(request: LoginRequest, auth: AuthDependency) -> AuthResponse:
    rate_limiter.check(f"auth:login:{request.email.lower()}", settings.auth_rate_limit_per_minute)
    user = auth.authenticate(request.email, request.password)
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid email or password.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    return AuthResponse(access_token=auth.create_token(user), user=UserResponse(**user.__dict__))


@app.post("/api/auth/forgot-password", response_model=MessageResponseBody)
async def forgot_password(
    request: ForgotPasswordRequest, auth: AuthDependency, config: SettingsDependency
) -> MessageResponseBody:
    rate_limiter.check(
        f"auth:forgot:{request.email.lower()}", config.auth_rate_limit_per_minute
    )
    if not config.email_configured:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=(
                "Password-reset email is not fully configured on this server. "
                "Add valid SMTP credentials and restart the backend."
            ),
        )
    token = auth.create_password_reset(request.email, config.password_reset_minutes)
    if token:
        reset_url = f"{config.frontend_public_url.rstrip('/')}?reset_token={token}"
        try:
            await send_password_reset_email(config, request.email.lower(), reset_url)
        except EmailConfigurationError as exc:
            raise HTTPException(status_code=503, detail=str(exc)) from exc
        except Exception as exc:
            logger.exception("Password-reset email delivery failed")
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="The reset email could not be sent. Please try again shortly.",
            ) from exc
    return MessageResponseBody(
        message="If an account exists for that email, a password-reset link has been sent."
    )


@app.post("/api/auth/reset-password", response_model=MessageResponseBody)
async def reset_password(
    request: ResetPasswordRequest, auth: AuthDependency
) -> MessageResponseBody:
    rate_limiter.check("auth:reset", settings.auth_rate_limit_per_minute)
    if not auth.reset_password(request.token, request.password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This password-reset link is invalid or has expired.",
        )
    return MessageResponseBody(message="Your password has been reset. You can now sign in.")


@app.get("/api/auth/me", response_model=UserResponse)
async def me(user: CurrentUser) -> UserResponse:
    return UserResponse(**user.__dict__)


@app.get("/api/boards", response_model=list[BoardSummaryResponse])
async def list_boards(user: CurrentUser, chats: ChatStoreDependency) -> list[dict]:
    return chats.list_boards(user.id)


@app.get("/api/search", response_model=list[SearchResultResponse])
async def search_messages(
    user: CurrentUser,
    chats: ChatStoreDependency,
    q: str = Query(min_length=2, max_length=100),
) -> list[dict]:
    query = q.strip()
    if len(query) < 2:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Search must contain at least 2 characters.",
        )
    return chats.search_messages(user.id, query)


@app.get("/api/media/{media_id}")
async def get_media(media_id: int, user: CurrentUser, chats: ChatStoreDependency) -> FileResponse:
    media = chats.get_media(user.id, media_id)
    return FileResponse(media["path"], media_type=media["mime_type"])


@app.post("/api/uploads", response_model=TemporaryUploadResponse, status_code=201)
async def upload_attachment(
    user: CurrentUser,
    chats: ChatStoreDependency,
    file: Annotated[UploadFile, File()],
) -> TemporaryUploadResponse:
    rate_limiter.check(f"upload:{user.id}", settings.upload_rate_limit_per_minute)
    filename = file.filename or ""
    if not filename or PurePath(filename).name != filename or "\x00" in filename:
        raise HTTPException(status_code=422, detail="Attachment filename is invalid.")
    extension = PurePath(filename).suffix.lower()
    expected_mime = SUPPORTED_ATTACHMENT_TYPES.get(extension)
    if expected_mime is None:
        raise HTTPException(status_code=422, detail=f"{filename} has an unsupported file format.")
    if file.content_type not in {expected_mime, "application/octet-stream", None, ""}:
        raise HTTPException(status_code=422, detail=f"{filename} has an invalid file type.")
    chats.cleanup_temporary_uploads()
    chats.temporary_upload_path.mkdir(parents=True, exist_ok=True)
    stored_filename = f"{uuid.uuid4().hex}{extension}"
    target = chats.temporary_upload_path / stored_filename
    size = 0
    try:
        with target.open("xb") as destination:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_BINARY_UPLOAD_BYTES:
                    raise AttachmentValidationError(f"{filename} is larger than 100 MB.")
                destination.write(chunk)
        if size == 0:
            raise AttachmentValidationError(f"{filename} is empty.")
        with target.open("rb") as uploaded_file:
            header = uploaded_file.read(12)
        valid_signature = {
            ".pdf": header.startswith(b"%PDF-"),
            ".jpg": header.startswith(b"\xff\xd8\xff"),
            ".jpeg": header.startswith(b"\xff\xd8\xff"),
            ".png": header.startswith(b"\x89PNG\r\n\x1a\n"),
            ".webp": header.startswith(b"RIFF") and header[8:12] == b"WEBP",
            ".docx": header.startswith(b"PK"),
            ".xlsx": header.startswith(b"PK"),
            ".pptx": header.startswith(b"PK"),
        }.get(extension, True)
        if not valid_signature:
            raise AttachmentValidationError(f"{filename} is corrupted or mislabeled.")
        if extension in {".docx", ".xlsx", ".pptx"}:
            _validate_office_archive(target, filename)
        if extension in IMAGE_ATTACHMENT_EXTENSIONS:
            try:
                from PIL import Image

                Image.MAX_IMAGE_PIXELS = 50_000_000
                with Image.open(target) as image:
                    image.verify()
            except Exception as exc:
                raise AttachmentValidationError(
                    f"{filename} is corrupted or exceeds the safe image limit."
                ) from exc
        current_storage = (
            chats.user_storage_bytes(user.id) if hasattr(chats, "user_storage_bytes") else 0
        )
        if current_storage + size > settings.max_user_storage_bytes:
            raise AttachmentValidationError(
                "Your stored attachments exceed the account storage limit. Delete an older chat "
                "or remove an upload and try again."
            )
        preview_type = (
            "pdf"
            if extension == ".pdf"
            else "image"
            if extension in IMAGE_ATTACHMENT_EXTENSIONS
            else "text"
            if extension in TEXT_ATTACHMENT_EXTENSIONS
            else "document"
        )
        result = chats.register_temporary_upload(
            user.id, filename, stored_filename, expected_mime, size, preview_type
        )
        return TemporaryUploadResponse(**result)
    except AttachmentValidationError as exc:
        target.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        await file.close()


@app.delete("/api/uploads/{upload_id}", status_code=204)
async def delete_temporary_upload(
    upload_id: str, user: CurrentUser, chats: ChatStoreDependency
) -> None:
    chats.delete_temporary_upload(user.id, upload_id)


@app.get("/api/attachments/{attachment_id}")
async def get_attachment(
    attachment_id: int, user: CurrentUser, chats: ChatStoreDependency
) -> FileResponse:
    attachment = chats.get_attachment(user.id, attachment_id)
    return FileResponse(
        attachment["path"],
        media_type=attachment["mime_type"],
        filename=attachment["filename"],
        content_disposition_type="inline",
    )


@app.get("/api/attachments/{attachment_id}/preview", response_model=AttachmentPreviewResponse)
async def get_attachment_preview(
    attachment_id: int, user: CurrentUser, chats: ChatStoreDependency
) -> AttachmentPreviewResponse:
    return AttachmentPreviewResponse(**chats.get_attachment_preview(user.id, attachment_id))


@app.post("/api/boards", response_model=BoardSummaryResponse, status_code=201)
async def create_board(
    request: BoardCreateRequest, user: CurrentUser, chats: ChatStoreDependency
) -> dict:
    return chats.create_board(user.id, request.title.strip())


@app.get("/api/boards/{board_id}", response_model=BoardDetailResponse)
async def get_board(board_id: int, user: CurrentUser, chats: ChatStoreDependency) -> dict:
    return chats.get_board(user.id, board_id)


@app.patch("/api/boards/{board_id}", response_model=BoardDetailResponse)
async def rename_board(
    board_id: int,
    request: BoardRenameRequest,
    user: CurrentUser,
    chats: ChatStoreDependency,
) -> dict:
    return chats.rename_board(user.id, board_id, request.title.strip())


@app.delete("/api/boards/{board_id}", status_code=204)
async def delete_board(board_id: int, user: CurrentUser, chats: ChatStoreDependency) -> None:
    chats.delete_board(user.id, board_id)


@app.post("/api/chat", response_model=ChatResponse)
async def chat(
    request: ChatRequest,
    service: Annotated[RAGService, Depends(get_rag_service)],
    user: CurrentUser,
    chats: ChatStoreDependency,
) -> ChatResponse:
    rate_limiter.check(f"chat:{user.id}", settings.chat_rate_limit_per_minute)
    try:
        attachment_parts, stored_attachments, pending_attachments = resolve_request_attachments(
            request, chats, user.id
        )
    except AttachmentValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=str(exc)
        ) from exc
    history = chats.recent_messages(user.id, request.board_id, limit=19)
    saved_question = request.question
    wants_image = is_image_request(request.question)
    wants_video = is_video_request(request.question)
    if wants_image or wants_video:
        user_message = chats.add_message(user.id, request.board_id, "user", saved_question)
        for attachment in stored_attachments:
            chats.add_attachment(user.id, user_message["id"], **attachment)
        for attachment in pending_attachments:
            chats.consume_temporary_upload(
                user.id, user_message["id"], attachment["id"], attachment["preview_text"]
            )
        started_at = time.monotonic()
        try:
            tasks = []
            if wants_image:
                tasks.append(("image", service.generate_image(request.question)))
            if wants_video:
                tasks.append(("video", service.generate_video(request.question)))
            async with asyncio.timeout(600):
                generated_items = await asyncio.gather(
                    *(task for _, task in tasks), return_exceptions=True
                )
            successful = [
                (media_type, generated)
                for (media_type, _), generated in zip(tasks, generated_items, strict=True)
                if not isinstance(generated, Exception)
            ]
            failures = [
                generated for generated in generated_items if isinstance(generated, Exception)
            ]
            if not successful:
                first_error = failures[0] if failures else None
                if isinstance(first_error, MediaGenerationError):
                    raise first_error
                raise MediaGenerationError(
                    "provider_error",
                    "OpenAI could not generate the requested media right now.",
                    True,
                )
            labels = [media_type for media_type, _ in successful]
            answer = (
                "Here are the image and video I generated for you."
                if len(labels) == 2
                else f"Here is the {labels[0]} I generated for you."
            )
            partial_error = next(
                (error for error in failures if isinstance(error, MediaGenerationError)), None
            )
            if partial_error:
                answer += f" {partial_error.safe_message}"
            message = chats.add_message(
                user.id,
                request.board_id,
                "assistant",
                answer,
                source="general",
                media_error=partial_error.payload() if partial_error else None,
            )
            media = [
                chats.add_media(
                    user.id,
                    message["id"],
                    media_type,
                    generated.mime_type,
                    generated.data,
                    generated.alt,
                )
                for media_type, generated in successful
            ]
            duration = time.monotonic() - started_at
            for item in media:
                logger.info(
                    "Generated media successfully: type=%s model=%s "
                    "duration_seconds=%.2f media_id=%s",
                    item["type"],
                    "gpt-image-2" if item["type"] == "image" else "sora-2",
                    duration,
                    item["id"],
                )
            return ChatResponse(
                answer=answer,
                citations=[],
                source="general",
                media=media,
                media_error=partial_error.payload() if partial_error else None,
            )
        except TimeoutError:
            logger.exception("OpenAI media generation timed out after ten minutes")
            error = MediaGenerationError(
                "generation_timeout",
                "Media generation timed out. Please try again.",
                True,
            )
            answer = error.safe_message
            chats.add_message(
                user.id,
                request.board_id,
                "assistant",
                answer,
                source="unavailable",
                media_error=error.payload(),
            )
            return ChatResponse(
                answer=answer,
                citations=[],
                source="unavailable",
                media=[],
                media_error=error.payload(),
            )
        except MediaGenerationError as error:
            answer = error.safe_message
            chats.add_message(
                user.id,
                request.board_id,
                "assistant",
                answer,
                source="unavailable",
                media_error=error.payload(),
            )
            return ChatResponse(
                answer=answer,
                citations=[],
                source="unavailable",
                media=[],
                media_error=error.payload(),
            )
        except RAGConfigurationError:
            error = MediaGenerationError(
                "missing_api_key", "Media generation is not configured on the server.", False
            )
            chats.add_message(
                user.id,
                request.board_id,
                "assistant",
                error.safe_message,
                source="unavailable",
                media_error=error.payload(),
            )
            return ChatResponse(
                answer=error.safe_message,
                citations=[],
                source="unavailable",
                media=[],
                media_error=error.payload(),
            )

    try:
        response = await service.ask(
            request.question, history=history, attachment_parts=attachment_parts
        )
        user_message = chats.add_message(user.id, request.board_id, "user", saved_question)
        for attachment in stored_attachments:
            chats.add_attachment(user.id, user_message["id"], **attachment)
        for attachment in pending_attachments:
            chats.consume_temporary_upload(
                user.id, user_message["id"], attachment["id"], attachment["preview_text"]
            )
        chats.add_message(
            user.id,
            request.board_id,
            "assistant",
            response.answer,
            [citation.model_dump() for citation in response.citations],
            response.source,
        )
        return response
    except RAGConfigurationError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc)
        ) from exc
    except RAGServiceError as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc


@app.post("/api/chat/stream")
async def stream_chat(
    request: ChatRequest,
    service: Annotated[RAGService, Depends(get_rag_service)],
    user: CurrentUser,
    chats: ChatStoreDependency,
) -> StreamingResponse:
    """Stream normal text answers as newline-delimited JSON events."""

    async def events():
        try:
            if is_image_request(request.question) or is_video_request(request.question):
                response = await chat(request, service, user, chats)
                yield json.dumps({"type": "complete", **response.model_dump()}) + "\n"
                return
            try:
                attachment_parts, stored_attachments, pending_attachments = (
                    resolve_request_attachments(request, chats, user.id)
                )
            except AttachmentValidationError as exc:
                yield json.dumps({"type": "error", "message": str(exc)}) + "\n"
                return
            history = chats.recent_messages(user.id, request.board_id, limit=10)
            completed: ChatResponse | None = None
            async for event in service.stream_answer(
                request.question,
                history=history,
                attachment_parts=attachment_parts,
            ):
                if event["type"] == "delta":
                    yield json.dumps(event) + "\n"
                else:
                    completed = event["response"]
            if completed is None:
                yield (
                    json.dumps(
                        {"type": "error", "message": "OpenAI did not complete the response."}
                    )
                    + "\n"
                )
                return
            user_message = chats.add_message(user.id, request.board_id, "user", request.question)
            for attachment in stored_attachments:
                chats.add_attachment(user.id, user_message["id"], **attachment)
            for attachment in pending_attachments:
                chats.consume_temporary_upload(
                    user.id,
                    user_message["id"],
                    attachment["id"],
                    attachment["preview_text"],
                )
            chats.add_message(
                user.id,
                request.board_id,
                "assistant",
                completed.answer,
                [citation.model_dump() for citation in completed.citations],
                completed.source,
            )
            yield json.dumps({"type": "complete", **completed.model_dump()}) + "\n"
        except (RAGConfigurationError, RAGServiceError) as exc:
            yield json.dumps({"type": "error", "message": str(exc)}) + "\n"
        except asyncio.CancelledError:
            raise
        except Exception:
            logger.exception("Unexpected streaming chat failure")
            yield (
                json.dumps(
                    {"type": "error", "message": "Could not generate an answer. Please try again."}
                )
                + "\n"
            )

    return StreamingResponse(
        events(),
        media_type="application/x-ndjson",
        headers={"Cache-Control": "no-cache", "X-Content-Type-Options": "nosniff"},
    )
