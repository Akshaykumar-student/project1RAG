import base64
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from backend.main import (
    AttachmentValidationError,
    _extract_office_text,
    _validate_office_archive,
    prepare_attachment_storage,
    prepare_attachments,
    prepare_temporary_uploads,
    resolve_request_attachments,
)
from backend.schemas import ChatAttachment, ChatRequest


def attachment(filename: str, mime_type: str, data: bytes, encoding="base64") -> ChatAttachment:
    content = data.decode() if encoding == "utf8" else base64.b64encode(data).decode()
    return ChatAttachment(
        filename=filename,
        mime_type=mime_type,
        encoding=encoding,
        size=len(data),
        content=content,
    )


def office_bytes(kind: str) -> bytes:
    stream = BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph("Document paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        document.save(stream)
    elif kind == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Name", "Value"])
        sheet.append(["Alpha", 42])
        workbook.save(stream)
        workbook.close()
    else:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Slide title"
        table = slide.shapes.add_table(1, 1, 0, 0, 100, 100).table
        table.cell(0, 0).text = "Cell"
        presentation.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    ("extension", "expected"),
    [(".docx", "Document paragraph"), (".xlsx", "Worksheet: Data"), (".pptx", "Slide title")],
)
def test_extracts_supported_office_text(extension, expected):
    text = _extract_office_text(extension, office_bytes(extension[1:]), f"file{extension}")

    assert expected in text


def test_validates_office_archives_and_rejects_bad_zip(tmp_path):
    valid = tmp_path / "valid.docx"
    valid.write_bytes(office_bytes("docx"))
    _validate_office_archive(valid, valid.name)

    invalid = tmp_path / "invalid.docx"
    invalid.write_bytes(b"not-a-zip")
    with pytest.raises(AttachmentValidationError, match="corrupted or mislabeled"):
        _validate_office_archive(invalid, invalid.name)


@pytest.mark.parametrize(
    ("item", "message"),
    [
        (attachment("../notes.txt", "text/plain", b"hello", "utf8"), "filenames"),
        (attachment("notes.exe", "application/octet-stream", b"hello"), "unsupported"),
        (attachment("notes.txt", "text/csv", b"hello", "utf8"), "invalid file type"),
        (attachment("bad.png", "image/png", b"not-png"), "corrupted or mislabeled"),
    ],
)
def test_attachment_validation_errors(item, message):
    with pytest.raises(AttachmentValidationError, match=message):
        prepare_attachments([item])


def test_rejects_malformed_base64_and_wrong_reported_size():
    malformed = ChatAttachment(
        filename="photo.png",
        mime_type="image/png",
        encoding="base64",
        size=5,
        content="not base64!",
    )
    with pytest.raises(AttachmentValidationError, match="invalid file data"):
        prepare_attachments([malformed])

    wrong_size = attachment("notes.txt", "text/plain", b"hello", "utf8")
    wrong_size.size = 99
    with pytest.raises(AttachmentValidationError, match="invalid reported size"):
        prepare_attachments([wrong_size])


def test_prepares_pdf_and_office_attachments():
    pdf_stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf_stream)
    pdf = attachment("file.pdf", "application/pdf", pdf_stream.getvalue())
    docx = attachment(
        "file.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        office_bytes("docx"),
    )

    parts = prepare_attachments([pdf, docx])

    assert parts[0]["type"] == "input_file"
    assert parts[1]["type"] == "input_text"
    assert "Document paragraph" in parts[1]["text"]


def test_prepares_persisted_attachment_metadata():
    text = attachment("notes.txt", "text/plain", b"hello", "utf8")
    image = attachment("photo.png", "image/png", b"\x89PNG\r\n\x1a\ncontent")
    document = attachment(
        "file.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        office_bytes("docx"),
    )

    stored = prepare_attachment_storage([text, image, document])

    assert [item["preview_type"] for item in stored] == ["text", "image", "document"]
    assert stored[0]["preview_text"] == "hello"
    assert "Document paragraph" in stored[2]["preview_text"]


def test_prepares_temporary_text_image_pdf_and_office_files(tmp_path):
    text_path = tmp_path / "notes.txt"
    text_path.write_text("temporary text", encoding="utf-8")
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (20, 20), "red").save(image_path)
    pdf_path = tmp_path / "file.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as output:
        writer.write(output)
    docx_path = tmp_path / "file.docx"
    docx_path.write_bytes(office_bytes("docx"))
    uploads = [
        {"id": path.stem, "filename": path.name, "path": path, "size": path.stat().st_size}
        for path in (text_path, image_path, pdf_path, docx_path)
    ]

    parts, prepared = prepare_temporary_uploads(uploads)

    assert [part["type"] for part in parts] == [
        "input_text",
        "input_image",
        "local_input_file",
        "input_text",
    ]
    assert prepared[0]["preview_text"] == "temporary text"
    assert "Document paragraph" in prepared[-1]["preview_text"]


def test_resolve_request_attachments_rejects_mixed_payloads():
    request = ChatRequest(
        board_id=1,
        question="Read these",
        attachment_ids=["upload-id"],
        attachments=[attachment("notes.txt", "text/plain", b"hello", "utf8")],
    )

    with pytest.raises(AttachmentValidationError, match="not both"):
        resolve_request_attachments(request, object(), 1)
